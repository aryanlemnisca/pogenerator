"""One-time script: copy PO-027 and clear variable cells to produce po_template.xlsx."""
from pathlib import Path
from openpyxl import load_workbook

SRC = Path(__file__).parent.parent.parent / "reference" / "PO-027_Vasa_Sci.xlsx"
DST = Path(__file__).parent.parent / "data" / "po_template.xlsx"

wb = load_workbook(SRC)
ws = wb.active

# Clear PO meta block
ws["A2"] = None
ws["E2"] = None

# Clear vendor block (rows 5-9, col A)
for row in range(5, 10):
    ws[f"A{row}"] = None

# Clear deliver-to block (rows 5-9, col E)
for row in range(5, 10):
    ws[f"E{row}"] = None

# Clear item rows 14-18: serial#, catalog, brand, description, qty, rate, dis%, gst%
# Leave I and K columns (formulas) — they'll be overwritten during generation
for row in range(14, 19):
    for col in ("A", "B", "C", "D", "F", "G", "H", "J"):
        ws[f"{col}{row}"] = None

# Clear footer rows 22-23
ws["A22"] = None
ws["A23"] = None

DST.parent.mkdir(parents=True, exist_ok=True)
wb.save(DST)
print(f"Template written to {DST}")
