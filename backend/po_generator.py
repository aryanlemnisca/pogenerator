"""
Template-based PO xlsx generator using openpyxl.

NOTE: The original PO-027 template has a known inconsistency — row 14 uses
=G14*(1-H14/100)*F14 (includes qty multiplier) while rows 15–18 use
=G{n}*(1-H{n}/100) (qty multiplier absent). We deliberately standardize all
rows to include *F{n} so Dis Rate always reflects the total discounted cost.
"""
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange

from schemas import POPayload

TEMPLATE_PATH = Path(__file__).parent / "data" / "po_template.xlsx"
DELIVER_TO_PATH = Path(__file__).parent / "data" / "deliver_to.json"

FIRST_ITEM_ROW = 14
TEMPLATE_ITEM_ROWS = 5  # PO-027 has rows 14–18


def _copy_cell_style(src_cell, dst_cell):
    """Copy font, fill, alignment, border, number_format from src to dst."""
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.border = copy(src_cell.border)
        dst_cell.number_format = src_cell.number_format


def _item_formulas(row: int) -> tuple[str, str]:
    """Return (dis_rate_formula, amt_formula) for a given spreadsheet row."""
    dis_rate = f"=G{row}*(1-H{row}/100)*F{row}"
    amt = f"=I{row}*(1+J{row}/100)"
    return dis_rate, amt


def generate_po(payload: POPayload, output_path: Path) -> None:
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb.active

    # 1. PO meta block (A2, merged A2:D2)
    meta = payload.meta
    ws["A2"] = (
        f"Purchase Order# : {meta.po_number.value}\n"
        f"Date : {meta.date.value}\n"
        f"Payment Terms : {meta.payment_terms.value}\n"
        f"Delivery Date : {meta.delivery_date.value}\n"
        f"Ref# : {meta.ref_number.value}"
    )

    # 2. Place-of-supply block (E2, merged E2:K2)
    ws["E2"] = (
        f"Place Of Supply : {meta.place_of_supply.value}\n"
        f"PO Revision : {meta.po_revision.value}\n"
        f"PO Type : {meta.po_type.value}\n"
        f"Inco Terms : {meta.inco_terms.value}\n"
        f"Dispatch Instructions : {meta.dispatch_instructions.value}"
    )

    # 3. Vendor block (A5–A9)
    vendor = payload.vendor
    ws["A5"] = vendor.name.value
    address_lines = [fl.value for fl in vendor.address_lines]
    ws["A6"] = address_lines[0] if len(address_lines) > 0 else ""
    ws["A7"] = address_lines[1] if len(address_lines) > 1 else ""
    ws["A8"] = f"GST: {vendor.gst.value}"
    ws["A9"] = vendor.contact_line.value

    # 4. Deliver-To block (E5–E9) — always fixed from JSON
    dt = json.loads(DELIVER_TO_PATH.read_text())
    ws["E5"] = dt["name"]
    ws["E6"] = dt["address_line1"]
    ws["E7"] = dt["address_line2"]
    ws["E8"] = dt["gstin"]
    ws["E9"] = dt["contact"]

    # 5. Resize item rows
    n_items = len(payload.items)
    last_item_row = FIRST_ITEM_ROW + n_items - 1

    if n_items > TEMPLATE_ITEM_ROWS:
        rows_to_add = n_items - TEMPLATE_ITEM_ROWS
        insert_after = FIRST_ITEM_ROW + TEMPLATE_ITEM_ROWS - 1
        insert_at = insert_after + 1  # first newly-inserted row number

        # openpyxl 3.x does not always shift merged-cell ranges that fall
        # entirely *below* the insertion point. Fix them manually before
        # calling insert_rows so we know the pre-insert coordinates.
        merged_to_fix = [
            str(mc) for mc in ws.merged_cells.ranges
            if mc.min_row >= insert_at
        ]
        for ref in merged_to_fix:
            ws.unmerge_cells(ref)

        ws.insert_rows(insert_at, rows_to_add)

        # Re-add the previously-unmerged ranges, shifted down by rows_to_add
        for ref in merged_to_fix:
            cr = CellRange(ref)
            new_ref = CellRange(
                min_col=cr.min_col,
                min_row=cr.min_row + rows_to_add,
                max_col=cr.max_col,
                max_row=cr.max_row + rows_to_add,
            )
            ws.merge_cells(str(new_ref))

        # Copy style from row 14 (style donor) into each new row
        donor_row = FIRST_ITEM_ROW
        for new_row in range(insert_at, insert_at + rows_to_add):
            for col_idx in range(1, 12):  # cols A–K
                src = ws.cell(row=donor_row, column=col_idx)
                dst = ws.cell(row=new_row, column=col_idx)
                _copy_cell_style(src, dst)
            # Re-merge D:E for description on the new row
            ws.merge_cells(f"D{new_row}:E{new_row}")

    elif n_items < TEMPLATE_ITEM_ROWS:
        rows_to_delete = TEMPLATE_ITEM_ROWS - n_items
        delete_start = FIRST_ITEM_ROW + n_items
        ws.delete_rows(delete_start, rows_to_delete)

    # 6. Write item data
    for idx, item in enumerate(payload.items):
        row = FIRST_ITEM_ROW + idx
        ws[f"A{row}"] = idx + 1
        ws[f"B{row}"] = item.catalog_number.value
        ws[f"C{row}"] = item.brand.value
        ws[f"D{row}"] = item.description.value
        ws[f"F{row}"] = item.qty.value
        ws[f"G{row}"] = item.rate.value
        ws[f"H{row}"] = item.discount_percent.value
        ws[f"J{row}"] = item.gst_percent.value
        dis_rate_formula, amt_formula = _item_formulas(row)
        ws[f"I{row}"] = dis_rate_formula
        ws[f"K{row}"] = amt_formula

    # 7. Total row (one blank gap row after last item)
    total_row = FIRST_ITEM_ROW + n_items + 1
    ws[f"K{total_row}"] = f"=SUM(K{FIRST_ITEM_ROW}:K{last_item_row})"

    # 8. Footer — Prepared By / Authorized Signature (shifts with item count)
    footer_base = total_row + 2
    ws[f"A{footer_base}"] = f"Prepared By: {payload.prepared_by}"
    ws[f"A{footer_base + 1}"] = f"Authorized Signature: {payload.authorized_signature}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
