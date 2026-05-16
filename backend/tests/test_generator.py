"""Tests for the xlsx generator. Run from the backend/ directory: pytest tests/ -v"""
import pytest
from pathlib import Path
from openpyxl import load_workbook
from schemas import POPayload, POMeta, VendorBlock, LineItem, FlaggableStr, FlaggableFloat
from po_generator import generate_po

OUTPUT_DIR = Path("tests/output")


def _make_vasa_payload() -> POPayload:
    return POPayload(
        meta=POMeta(
            po_number=FlaggableStr(value="PO-027", was_found=False),
            date=FlaggableStr(value="15 May 2026", was_found=False),
            payment_terms=FlaggableStr(value="30 days", was_found=False),
            delivery_date=FlaggableStr(value="30 May 2026", was_found=True),
            ref_number=FlaggableStr(value="2603122", was_found=True),
            place_of_supply=FlaggableStr(value="Karnataka", was_found=False),
        ),
        vendor=VendorBlock(
            name=FlaggableStr(value="Vasa Scientific Co.", was_found=True),
            address_lines=[
                FlaggableStr(value="95/1, 11th cross, 4th main, Malleshwaram", was_found=True),
                FlaggableStr(value="Bangalore, 560003", was_found=True),
            ],
            gst=FlaggableStr(value="29AAMPP7088B1Z8", was_found=True),
            contact_line=FlaggableStr(
                value="Telephone: (080) 43023447; M: 9611271358; enquiry@vasascientific.in",
                was_found=True,
            ),
        ),
        items=[
            LineItem(
                catalog_number=FlaggableStr(value="T460091", was_found=True),
                brand=FlaggableStr(value="Tarsons", was_found=True),
                description=FlaggableStr(value="PETRI DISH 90MM (6 x 1 = 6 NOS)", was_found=True),
                qty=FlaggableFloat(value=2, was_found=True),
                rate=FlaggableFloat(value=5592, was_found=True),
                discount_percent=FlaggableFloat(value=18, was_found=True),
                gst_percent=FlaggableFloat(value=18, was_found=True),
            ),
            LineItem(
                catalog_number=FlaggableStr(value="T560090", was_found=True),
                brand=FlaggableStr(value="Tarsons", was_found=True),
                description=FlaggableStr(value="WASH BOTTLE LDPE 500ML (SLEEVE OF 10)", was_found=True),
                qty=FlaggableFloat(value=1, was_found=True),
                rate=FlaggableFloat(value=1095, was_found=True),
                discount_percent=FlaggableFloat(value=18, was_found=True),
                gst_percent=FlaggableFloat(value=18, was_found=True),
            ),
            LineItem(
                catalog_number=FlaggableStr(value="T630090", was_found=True),
                brand=FlaggableStr(value="Tarsons", was_found=True),
                description=FlaggableStr(value="FUNNEL PP 160MM", was_found=True),
                qty=FlaggableFloat(value=1, was_found=True),
                rate=FlaggableFloat(value=966, was_found=True),
                discount_percent=FlaggableFloat(value=18, was_found=True),
                gst_percent=FlaggableFloat(value=18, was_found=True),
            ),
            LineItem(
                catalog_number=FlaggableStr(value="T202080", was_found=True),
                brand=FlaggableStr(value="Tarsons", was_found=True),
                description=FlaggableStr(value="UNIVERSAL COMBI RACK", was_found=True),
                qty=FlaggableFloat(value=1, was_found=True),
                rate=FlaggableFloat(value=3459, was_found=True),
                discount_percent=FlaggableFloat(value=18, was_found=True),
                gst_percent=FlaggableFloat(value=18, was_found=True),
            ),
            LineItem(
                catalog_number=FlaggableStr(value="T720510", was_found=True),
                brand=FlaggableStr(value="Tarsons", was_found=True),
                description=FlaggableStr(value="PINCH CLAMP PP", was_found=True),
                qty=FlaggableFloat(value=1, was_found=True),
                rate=FlaggableFloat(value=1027, was_found=True),
                discount_percent=FlaggableFloat(value=18, was_found=True),
                gst_percent=FlaggableFloat(value=18, was_found=True),
            ),
        ],
    )


@pytest.fixture(autouse=True)
def output_dir():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def test_generate_po_produces_file():
    payload = _make_vasa_payload()
    out = OUTPUT_DIR / "test_po_basic.xlsx"
    generate_po(payload, out)
    assert out.exists()
    assert out.stat().st_size > 10_000


def test_meta_block_written():
    payload = _make_vasa_payload()
    out = OUTPUT_DIR / "test_po_meta.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    a2 = ws["A2"].value
    assert "PO-027" in a2
    assert "2603122" in a2
    assert "30 days" in a2


def test_vendor_block_written():
    payload = _make_vasa_payload()
    out = OUTPUT_DIR / "test_po_vendor.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    assert ws["A5"].value == "Vasa Scientific Co."
    assert "29AAMPP7088B1Z8" in ws["A8"].value


def test_deliver_to_block_written():
    payload = _make_vasa_payload()
    out = OUTPUT_DIR / "test_po_deliver.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    assert "BIOAI Innovations" in ws["E5"].value


def test_item_rows_written_with_formulas():
    payload = _make_vasa_payload()
    out = OUTPUT_DIR / "test_po_items.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    assert ws["B14"].value == "T460091"
    assert ws["G14"].value == 5592
    assert ws["H14"].value == 18
    assert "G14" in ws["I14"].value
    assert "F14" in ws["I14"].value   # must include qty multiplier
    assert "I14" in ws["K14"].value
    assert ws["B18"].value == "T720510"
    assert "F18" in ws["I18"].value   # all rows must have qty multiplier


def test_total_formula_correct():
    payload = _make_vasa_payload()
    out = OUTPUT_DIR / "test_po_total.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    total_cell = ws["K20"].value
    assert total_cell is not None
    assert "SUM(K14:K18)" in total_cell


def test_footer_written():
    payload = _make_vasa_payload()
    out = OUTPUT_DIR / "test_po_footer.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    assert "Leelakrishna" in ws["A22"].value
    assert "Shilpa Nargund" in ws["A23"].value


def test_extra_items_inserts_rows():
    """Seven-item PO: item rows 14–20, total at row 22."""
    payload = _make_vasa_payload()
    extra_item = LineItem(
        catalog_number=FlaggableStr(value="EXTRA", was_found=True),
        brand=FlaggableStr(value="TestBrand", was_found=True),
        description=FlaggableStr(value="Extra item", was_found=True),
        qty=FlaggableFloat(value=1, was_found=True),
        rate=FlaggableFloat(value=100, was_found=True),
        discount_percent=FlaggableFloat(value=0, was_found=True),
        gst_percent=FlaggableFloat(value=18, was_found=True),
    )
    payload.items.extend([extra_item, extra_item])  # now 7 items
    out = OUTPUT_DIR / "test_po_7items.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    assert ws["A20"].value == 7
    assert "SUM(K14:K20)" in ws["K22"].value


def test_fewer_items_deletes_rows():
    """Two-item PO: total at row 17 (items 14-15, blank 16, total 17)."""
    payload = _make_vasa_payload()
    payload.items = payload.items[:2]
    out = OUTPUT_DIR / "test_po_2items.xlsx"
    generate_po(payload, out)
    wb = load_workbook(out)
    ws = wb.active
    assert ws["A15"].value == 2
    assert "SUM(K14:K15)" in ws["K17"].value
